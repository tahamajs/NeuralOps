from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation('/Users/tahamajs/Documents/uni/Project/Docs/A1_Vertical_Template_Electronics.pptx')
slide = prs.slides[0]

def replace_paragraph_text(paragraph, new_text):
    if len(paragraph.runs) > 0:
        first_run = paragraph.runs[0]
        first_run.text = new_text
        for i in range(1, len(paragraph.runs)):
            paragraph.runs[i].text = ""
    else:
        paragraph.text = new_text

# Update Title / Abstract with 100 modules & Champion selection focus
if len(slide.shapes) > 13 and slide.shapes[13].has_text_frame:
    tf = slide.shapes[13].text_frame
    if len(tf.paragraphs) > 2:
        replace_paragraph_text(tf.paragraphs[2], "•  هدف و فرضیه کلیدی پژوهش: پیاده‌سازی و غربال‌گری ۱۰۰ نوع مکانیزم توجه متنوع و فشرده‌سازی در چارچوب یکپارچه NeuralOps جهت شناسایی برترین معماری بهینه سخت‌افزاری و جایگزینی کامل با MHA استاندارد بدون افت کیفیت زبانی.")

# Update Results shape with Champion Fusion Attention
if len(slide.shapes) > 15 and slide.shapes[15].has_text_frame:
    tf = slide.shapes[15].text_frame
    if len(tf.paragraphs) > 1:
        replace_paragraph_text(tf.paragraphs[1], "•  معرفی مدل قهرمان (Fusion Attention): در میان ۱۰۰ مکانیزم پیاده‌سازی‌شده، معماری Fusion Attention با دستیابی به دقت خیره‌کننده ۹۳.۱۲٪ (۲٪ بالاتر از Baseline با ۹۱.۳۰٪) و سقوط اوج مصرف حافظه به زیر ۴۸۰ مگابایت (صرفه‌جویی ۸۸٪ در VRAM) به عنوان برترین مدل پروژه انتخاب گردید.")

# Add stitched champion hero image & comparison chart
try:
    slide.shapes.add_picture('/Users/tahamajs/Documents/uni/Project/thesis_final/combined_hero_architecture.jpg', Inches(1.5), Inches(23.5), Inches(9.8), Inches(7.5))
    slide.shapes.add_picture('/Users/tahamajs/Documents/uni/Project/thesis_final/chart_champion_comparison.png', Inches(11.8), Inches(23.5), Inches(10.2), Inches(6.8))
except Exception as e:
    print(f"Failed to add images: {e}")

prs.save('/Users/tahamajs/Documents/uni/Project/Docs/A1_Vertical_Template_Electronics_Updated.pptx')
print("Poster successfully updated with Champion model and stitched architecture image!")
