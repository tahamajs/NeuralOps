from pptx import Presentation
prs = Presentation('/Users/tahamajs/Documents/uni/Project/Docs/A1_Vertical_Template_Electronics.pptx')
slide = prs.slides[0]
for i in [13, 14, 15, 17, 19, 20]:
    shape = slide.shapes[i]
    print(f"\n--- Shape {i} ---")
    print(shape.text_frame.text)
